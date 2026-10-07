import os
from fastmcp import FastMCP
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd

# ==================== 输出目录配置（基于 __file__ 解析，不依赖 CWD） ====================
CREAT_FILE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "creat_file")
os.makedirs(CREAT_FILE_DIR, exist_ok=True)

# 初始化MCP服务
mcp = FastMCP("DocumentTools")

# 工具内部辅助函数：安全拼接到creat_file目录，只保留文件名，防止传入完整路径越权
def get_target_filepath(input_file_path: str):
    # 只提取文件名，丢弃传入的任何目录部分
    filename = os.path.basename(input_file_path)
    return os.path.join(CREAT_FILE_DIR, filename)

# ===================== DOCX 工具 =====================
@mcp.tool()
def read_docx(file_path: str) -> str:
    """读取creat_file目录下docx文档全部文本内容（包含表格内文字）"""
    file_path = get_target_filepath(file_path)
    if not os.path.exists(file_path):
        return f"错误：文件不存在 {file_path}"
    doc = Document(file_path)
    content_lines = []
    for para in doc.paragraphs:
        if para.text.strip():
            content_lines.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            content_lines.append(" | ".join(row_cells))
    return "\n".join(content_lines)

@mcp.tool()
def create_docx(file_path: str, title: str, content: list[str]) -> str:
    """创建全新docx文档，写入标题与多段正文。生成的文件保存到 creat_file 目录。"""
    file_path = get_target_filepath(file_path)
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for para_text in content:
        doc.add_paragraph(para_text)
    doc.save(file_path)
    return f"成功创建docx，路径：{file_path}"

@mcp.tool()
def add_table_to_docx(file_path: str, table_data: list[list[str]]) -> str:
    """在creat_file目录下现有docx文档末尾插入表格。"""
    file_path = get_target_filepath(file_path)
    if not os.path.exists(file_path):
        return f"错误：文件不存在 {file_path}"
    doc = Document(file_path)
    if not table_data:
        return "表格数据不能为空"
    row_count = len(table_data)
    col_count = len(table_data[0])
    table = doc.add_table(rows=row_count, cols=col_count)
    table.style = 'Table Grid'
    for row_idx, row_data in enumerate(table_data):
        for col_idx, cell_text in enumerate(row_data):
            table.cell(row_idx, col_idx).text = cell_text
    doc.save(file_path)
    return f"成功向docx插入{row_count}行{col_count}列表格，路径：{file_path}"

# ===================== Excel表格工具 =====================
@mcp.tool()
def read_excel_table(file_path: str, sheet_name: str = "Sheet1") -> str:
    """读取creat_file目录下Excel表格，转为文本格式返回"""
    file_path = get_target_filepath(file_path)
    if not os.path.exists(file_path):
        return f"错误：文件不存在 {file_path}"
    df = pd.read_excel(file_path, sheet_name=sheet_name)
    return df.to_string()

@mcp.tool()
def write_excel_table(file_path: str, table_data: list[list[str]], sheet_name: str = "Sheet1") -> str:
    """将二维表格数据写入Excel文件。生成的文件保存到 creat_file 目录。"""
    file_path = get_target_filepath(file_path)
    headers = table_data[0]
    rows = table_data[1:]
    df = pd.DataFrame(rows, columns=headers)
    df.to_excel(file_path, sheet_name=sheet_name, index=False)
    return f"成功写入Excel表格至 {file_path}"

if __name__ == "__main__":
    # 端口8890，0.0.0.0允许局域网访问
    mcp.run(transport="http", host="0.0.0.0", port=8890)
